# Owner: Dev B
"""Batch upload, column mapping, validation and commit."""

from __future__ import annotations

import csv
import difflib
import io
import re
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE_PATH = Path(__file__).resolve().parents[3] / "shared" / "csv" / "admissions-template.csv"
# Fallback copy for deployments that ship backend/ without shared/. Keep in sync.
_FALLBACK_COLUMNS = (
    "institution_aicte_code,programme_code,academic_year,student_name,dob,gender,category,pwd,"
    "apaar_id,apaar_opt_out,provisional_id,nationality,oci_status,fn_status,nri_category,"
    "seat_pool,admission_mode,admission_date,qualifying_exam,qualifying_exam_pct,"
    "domicile_state,contact_phone,contact_email,guardian_name"
)
TEMPLATE_COLUMNS: list[str] = (
    (TEMPLATE_PATH.read_text(encoding="utf-8") if TEMPLATE_PATH.exists() else _FALLBACK_COLUMNS)
    .strip()
    .split(",")
)


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_")


def read_header(file_name: str, content: bytes) -> list[str]:
    """Return the header row of an uploaded CSV or XLSX file."""
    if file_name.lower().endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(content), read_only=True).active
        first = next(sheet.iter_rows(max_row=1, values_only=True), ())
        return [str(c) for c in first if c is not None]
    text = content.decode("utf-8-sig", errors="replace")
    return next(csv.reader(io.StringIO(text)), [])


def suggest_mapping(headers: list[str]) -> tuple[dict[str, str], list[str]]:
    """Map uploaded column names to template columns by normalised name, then fuzzy match."""
    mapping: dict[str, str] = {}
    unmapped: list[str] = []
    for header in headers:
        norm = _norm(header)
        match = norm if norm in TEMPLATE_COLUMNS else None
        if match is None:
            close = difflib.get_close_matches(norm, TEMPLATE_COLUMNS, n=1, cutoff=0.75)
            match = close[0] if close else None
        if match and match not in mapping.values():
            mapping[header] = match
        else:
            unmapped.append(header)
    return mapping, unmapped


def validate_batch(batch_id: str) -> None:
    """Apply every rule in shared/rule-codes.md to the mapped rows and store issues."""
    raise NotImplementedError


def commit_batch(batch_id: str) -> None:
    """Consume Seat Ledger seats, write admissions, open review cases on overrun."""
    raise NotImplementedError
